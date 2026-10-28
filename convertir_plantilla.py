"""
CONVERSIÓN A LA PLANTILLA DE IMPORTACIÓN

Toma un archivo "sucio" (p. ej. exportación de formularios) y genera:
  1. Un Excel con la estructura EXACTA de la plantilla (solo la hoja de datos).
  2. Un Excel aparte con el reporte de auditoría.

Reglas por columna destino (campo "regla" del mapeo):
  texto, correo, numero_documento, telefono_codigo, telefono_numero,
  nivel, documento, programa, periodo

Principios: no se inventan valores, las equivalencias solo se aplican si están
confirmadas en la configuración y cualquier dato dudoso se conserva y se marca
(celda resaltada + fila en el reporte).

Uso:
    python convertir_plantilla.py --entrada sucio.xlsx --plantilla plantilla.xlsx \
        --config config_aspirantes.json
"""

from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

import automatizacion_excel as ae

REGLAS_CATALOGO = {"nivel": "col_nivel", "documento": "col_documento", "programa": "col_programa"}
RELLENO_REVISION = PatternFill("solid", fgColor="FFF1D6")
RELLENO_INVALIDO = PatternFill("solid", fgColor="FBD9D0")
RE_CORREO = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


# ---------------------------- reglas por valor ----------------------------

def limpiar_texto(valor):
    """Quita espacios sobrantes. Devuelve (valor_final, cambio)."""
    if ae.es_vacio(valor):
        return None, False
    texto = re.sub(r"\s+", " ", str(valor).replace("\u00a0", " ")).strip()
    return texto, texto != str(valor)


def texto_numero_documento(valor):
    if ae.es_vacio(valor):
        return None
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return re.sub(r"\s+", "", str(valor)) if not isinstance(valor, int) else str(valor)


def separar_telefono(valor):
    """Devuelve (codigo, numero, estado, motivo). Estados: VACIO, CORREGIDO, VALIDO, REQUIERE_REVISION."""
    if ae.es_vacio(valor):
        return None, None, "VACIO", ""
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return None, str(int(valor)), "VALIDO", ""
    texto = re.sub(r"\s+", " ", str(valor).replace("\u00a0", " ")).strip()

    nacional, codigo, modificado, motivo = ae.normalizar_telefono(texto)
    if modificado:
        return codigo, nacional, "CORREGIDO", "Indicativo +57 separado del número."
    if motivo:
        return None, texto, "REQUIERE_REVISION", motivo
    if texto.startswith("+"):
        return None, texto, "REQUIERE_REVISION", (
            "Indicativo distinto de +57: no se separa automáticamente (longitud del indicativo desconocida)."
        )
    return None, texto, "VALIDO", ""


def validar_periodo(valor, cfg):
    """Devuelve (valor_final, estado, motivo) aplicando formato, sufijos y periodos permitidos."""
    if ae.es_vacio(valor):
        return None, "VACIO", ""

    candidato, nota = valor, ""
    if isinstance(valor, str) and cfg.periodo_sufijos:
        m = re.fullmatch(r"\s*(\d{4})\s*-\s*(\d)\s*", valor)
        if m and m.group(2) in cfg.periodo_sufijos:
            candidato = f"{m.group(1)}{cfg.periodo_sufijos[m.group(2)]}"
            nota = f"«{valor.strip()}» equivale a {candidato}. "

    final, estado, motivo = ae.normalizar_periodo(candidato, cfg.periodos, cfg.anio_min, cfg.anio_max)
    if estado in ("CORREGIDO", "VALIDO"):
        codigo = final[:6]
        if cfg.periodos_permitidos and codigo not in cfg.periodos_permitidos:
            return valor, "REQUIERE_REVISION", (
                f"{nota}El periodo {codigo} no está entre los permitidos "
                f"({', '.join(cfg.periodos_permitidos)})."
            )
        if estado == "CORREGIDO":
            return final, "CORREGIDO", (nota + motivo).strip()
    if estado == "REQUIERE_REVISION":
        return valor, estado, (nota + motivo).strip()
    return final, estado, motivo


# ---------------------------- conversión principal ----------------------------

def _indice_columnas(hoja):
    """Encabezado normalizado -> lista de columnas (para detectar duplicados)."""
    idx = {}
    for celda in hoja[1]:
        if not ae.es_vacio(celda.value):
            idx.setdefault(ae.normalizar_encabezado(celda.value), []).append(celda.column)
    return idx


def convertir_a_plantilla(entrada, plantilla, salida, salida_reporte, cfg):
    entrada, plantilla = Path(entrada), Path(plantilla)
    salida, salida_reporte = Path(salida), Path(salida_reporte)
    for ruta in (entrada, plantilla):
        if not ruta.exists():
            raise FileNotFoundError(f"No se encontró el archivo: {ruta}")
    if not cfg.mapeo:
        raise ValueError("La configuración no tiene un mapeo de columnas definido.")
    if len({entrada.resolve(), salida.resolve(), salida_reporte.resolve(), plantilla.resolve()}) < 4:
        raise ValueError("Los archivos de entrada, plantilla y salida deben ser distintos.")

    origen = load_workbook(entrada, data_only=False).worksheets[0]
    libro = load_workbook(plantilla)
    destino = libro.worksheets[0]

    col_origen = _indice_columnas(origen)
    col_destino = _indice_columnas(destino)

    aud = ae.Auditoria(cfg.registrar_validos)
    reg = aud.registrar
    hoja_nombre = origen.title
    avisos = []

    # ---- Resolver el mapeo (columna origen / columna destino) ----
    reglas = []
    for m in cfg.mapeo:
        d = col_destino.get(ae.normalizar_encabezado(m["destino"]))
        if not d or len(d) != 1:
            raise ValueError(f"La plantilla no tiene (o repite) la columna destino «{m['destino']}».")
        pos = None
        if m.get("origen"):
            o = col_origen.get(ae.normalizar_encabezado(m["origen"]), [])
            if len(o) > 1:
                reg(hoja_nombre, "", m["destino"], "", "", "REQUIERE_REVISION",
                    f"Encabezado de origen «{m['origen']}» duplicado; no se procesa.")
            elif not o:
                reg(hoja_nombre, "", m["destino"], "", "", "COLUMNA_NO_ENCONTRADA",
                    f"No se encontró la columna de origen «{m['origen']}».")
            else:
                pos = o[0]
        reglas.append({**m, "col_dest": d[0], "col_orig": pos})
    por_regla = {r["regla"]: r for r in reglas}

    # ---- Limpiar el ejemplo de la plantilla ----
    if destino.max_row >= 2:
        destino.delete_rows(2, destino.max_row)

    # ---- Filas ----
    fila_out = 1
    filas_origen = 0
    for fila in range(2, origen.max_row + 1):
        valores_fila = [c.value for c in origen[fila]]
        if all(ae.es_vacio(v) for v in valores_fila):
            continue
        filas_origen += 1
        fila_out += 1

        def poner(col, valor, estado=None, columna="", original=None, motivo=""):
            celda = destino.cell(row=fila_out, column=col)
            celda.value = valor
            if isinstance(valor, str):
                celda.number_format = "@"
                celda.data_type = "s"
            if estado == "REQUIERE_REVISION":
                celda.fill = RELLENO_REVISION
            elif estado == "INVALIDO":
                celda.fill = RELLENO_INVALIDO

        for r in reglas:
            regla, nombre = r["regla"], r["destino"]
            if regla == "telefono_codigo":
                continue  # se resuelve junto con telefono_numero
            crudo = origen.cell(row=fila, column=r["col_orig"]).value if r["col_orig"] else None

            if regla == "telefono_numero":
                codigo, numero, estado, motivo = separar_telefono(crudo)
                rc = por_regla.get("telefono_codigo")
                poner(r["col_dest"], numero, estado)
                if rc:
                    poner(rc["col_dest"], codigo, estado)
                if estado == "CORREGIDO":
                    reg(hoja_nombre, fila, nombre, crudo, numero, estado, motivo)
                    if rc:
                        reg(hoja_nombre, fila, rc["destino"], "", codigo, "CORREGIDO",
                            "Código de país tomado del teléfono.")
                elif estado == "REQUIERE_REVISION":
                    reg(hoja_nombre, fila, nombre, crudo, numero, estado, motivo)
                else:
                    aud.contar(hoja_nombre, nombre, estado)
                continue

            if regla in REGLAS_CATALOGO:
                clave = getattr(cfg, REGLAS_CATALOGO[regla])
                limpio, _ = limpiar_texto(crudo)
                final, estado, motivo = ae.validar_valor(
                    limpio, cfg.valores_validos.get(clave, []), cfg.equivalencias.get(clave, {}))
                poner(r["col_dest"], final if estado != "VACIO" else None, estado)
                if estado == "SIN_CATALOGO":
                    aud.sin_catalogo[(hoja_nombre, nombre)] += 1
                    aud.contar(hoja_nombre, nombre, estado)
                elif estado != "VACIO":
                    reg(hoja_nombre, fila, nombre, crudo, final, estado, motivo)
                continue

            if regla == "periodo":
                final, estado, motivo = validar_periodo(crudo, cfg)
                poner(r["col_dest"], final, estado)
                if estado != "VACIO":
                    reg(hoja_nombre, fila, nombre, crudo, final, estado, motivo)
                continue

            if regla == "numero_documento":
                final = texto_numero_documento(crudo)
                poner(r["col_dest"], final)
                if final is not None and str(crudo) != final:
                    reg(hoja_nombre, fila, nombre, crudo, final, "CORREGIDO",
                        "Número convertido a texto para conservar el formato.")
                continue

            if regla == "correo":
                final, cambio = limpiar_texto(crudo)
                if final is not None and not RE_CORREO.match(final):
                    poner(r["col_dest"], final, "REQUIERE_REVISION")
                    reg(hoja_nombre, fila, nombre, crudo, final, "REQUIERE_REVISION",
                        "El correo no tiene un formato válido.")
                else:
                    poner(r["col_dest"], final)
                    if cambio:
                        reg(hoja_nombre, fila, nombre, crudo, final, "CORREGIDO", "Espacios sobrantes eliminados.")
                continue

            # texto libre
            final, cambio = limpiar_texto(crudo)
            poner(r["col_dest"], final)
            if cambio:
                reg(hoja_nombre, fila, nombre, crudo, final, "CORREGIDO", "Espacios sobrantes eliminados.")

    # ---- Resumen de columnas sin catálogo ----
    for (h, c), n in sorted(aud.sin_catalogo.items()):
        aud.filas.append({
            "Hoja": h, "Fila": "", "Columna": c, "Valor original": f"{n} celda(s) con valor",
            "Valor final": "", "Resultado": "SIN_CATALOGO",
            "Motivo": "Catálogo oficial pendiente: los valores se conservaron sin cambios.",
        })

    libro.save(salida)

    rep = Workbook()
    ae._escribir_reportes(rep, aud)
    if "Sheet" in rep.sheetnames:
        del rep["Sheet"]
    rep.save(salida_reporte)

    resumen = Counter()
    for (_, _, estado), n in aud.contador.items():
        resumen[estado] += n
    return {
        "salida": salida, "salida_reporte": salida_reporte,
        "filas_origen": filas_origen, "filas_salida": fila_out - 1,
        "reporte": aud.filas, "resumen": dict(resumen), "avisos": avisos,
    }


def main(argv=None):
    p = argparse.ArgumentParser(description="Convierte un Excel a la plantilla de importación.")
    p.add_argument("--entrada", required=True)
    p.add_argument("--plantilla", required=True)
    p.add_argument("--config", required=True)
    p.add_argument("--salida", default="importacion_lista.xlsx")
    p.add_argument("--reporte", default="importacion_reporte.xlsx")
    a = p.parse_args(argv)
    cfg = ae.cargar_configuracion_json(a.config)
    r = convertir_a_plantilla(a.entrada, a.plantilla, a.salida, a.reporte, cfg)
    print(f"Filas leídas: {r['filas_origen']}  |  filas generadas: {r['filas_salida']}")
    for k, v in sorted(r["resumen"].items()):
        print(f"  {k:<22}{v}")
    print(f"Archivo para importar: {r['salida']}\nReporte: {r['salida_reporte']}")


if __name__ == "__main__":
    main()
