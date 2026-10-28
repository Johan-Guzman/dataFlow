"""
AUTOMATIZACIÓN DE EXCEL

Reglas aplicadas (solo sobre las columnas configuradas):
1. Nivel educativo, tipo de documento y programa de interés: se ajustan
   al catálogo oficial configurado (sin coincidencias aproximadas).
2. categoria_servicio: "_" -> espacios y formato Title Case.
3. Separación del indicativo telefónico +57.
4. Normalización de la columna periodo.
5. Hojas "Reporte_validacion" y "Resumen_validacion" con la auditoría.

El archivo original nunca se sobrescribe y no se eliminan filas.

Uso:
    python automatizacion_excel.py
    python automatizacion_excel.py --entrada datos_originales.xlsx \
        --salida datos_corregidos.xlsx --config config.json
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import Counter
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

# ============================================================
# 1. CONFIGURACIÓN CENTRALIZADA  (editar aquí)
# ============================================================

ARCHIVO_ENTRADA = Path("datos_originales.xlsx")
ARCHIVO_SALIDA = Path("datos_corregidos.xlsx")

# Encabezados de las columnas a procesar (se comparan ignorando
# mayúsculas, tildes, espacios y guiones bajos).
COLUMNA_NIVEL_EDUCATIVO = "nivel educativo"
COLUMNA_TIPO_DOCUMENTO = "tipo de documento"
COLUMNA_PROGRAMA_INTERES = "programa de interés"
COLUMNA_TEXTO_NORMALIZAR = "categoria_servicio"
COLUMNA_TELEFONO = "telefono"
COLUMNA_CODIGO_PAIS = "Codigo_Pais"
COLUMNA_PERIODO = "periodo"

HOJA_REPORTE = "Reporte_validacion"
HOJA_RESUMEN = "Resumen_validacion"
HOJA_REGLAS = "Reglas de validación"

# ------------------------------------------------------------
# CATÁLOGOS OFICIALES (PENDIENTES DE DEFINIR)
# ------------------------------------------------------------
# Escribe cada valor exactamente como debe quedar en el Excel final.
# Mientras una lista esté vacía NO se hacen correcciones categóricas
# para esa columna: los valores se conservan y se registran como
# SIN_CATALOGO (una sola fila de resumen por hoja y columna).

VALORES_VALIDOS = {
    COLUMNA_NIVEL_EDUCATIVO: [
        # "Bachillerato", "Técnico", "Tecnólogo", "Profesional",
    ],
    COLUMNA_TIPO_DOCUMENTO: [
        # "Cédula de ciudadanía", "Tarjeta de identidad", ...
    ],
    COLUMNA_PROGRAMA_INTERES: [
        # Nombres oficiales de los programas.
    ],
}

# ------------------------------------------------------------
# EQUIVALENCIAS CONFIRMADAS (clave -> valor oficial)
# ------------------------------------------------------------
# Solo se aplican las que estén aquí. Ninguna abreviatura se
# interpreta por suposición.

EQUIVALENCIAS = {
    COLUMNA_NIVEL_EDUCATIVO: {},
    COLUMNA_TIPO_DOCUMENTO: {
        # "CC": "Cédula de ciudadanía",
        # "TI": "Tarjeta de identidad",
        # "CE": "Cédula de extranjería",
    },
    COLUMNA_PROGRAMA_INTERES: {
        # Solo para variantes que NO se resuelven ignorando
        # tildes/mayúsculas. Ej.: "Ing. de Sistemas": "Ingeniería de Sistemas"
    },
}

# ------------------------------------------------------------
# PERIODOS ACADÉMICOS
# ------------------------------------------------------------

PERIODOS_ACADEMICOS = {
    "10": "Primer Semestre",
    "20": "Segundo Semestre",
}
PERIODO_ANIO_MINIMO = 2000
PERIODO_ANIO_MAXIMO = 2100

# ------------------------------------------------------------
# OPCIONES
# ------------------------------------------------------------

# Si es True y existe la hoja "Reglas de validación", lee los
# catálogos desde ella (solo columnas cuyo encabezado coincida
# exactamente con las columnas configuradas). Desactivado por defecto
# porque su estructura aún no se ha inspeccionado.
LEER_CATALOGOS_DESDE_HOJA = False

# Si es True, el reporte incluye una fila por cada celda VALIDA.
# Por defecto solo se cuentan en la hoja de resumen.
REGISTRAR_VALIDOS_EN_REPORTE = False


@dataclass
class Configuracion:
    col_nivel: str = COLUMNA_NIVEL_EDUCATIVO
    col_documento: str = COLUMNA_TIPO_DOCUMENTO
    col_programa: str = COLUMNA_PROGRAMA_INTERES
    col_texto: str = COLUMNA_TEXTO_NORMALIZAR
    col_telefono: str = COLUMNA_TELEFONO
    col_codigo_pais: str = COLUMNA_CODIGO_PAIS
    col_periodo: str = COLUMNA_PERIODO
    valores_validos: dict = field(
        default_factory=lambda: {k: list(v) for k, v in VALORES_VALIDOS.items()}
    )
    equivalencias: dict = field(
        default_factory=lambda: {k: dict(v) for k, v in EQUIVALENCIAS.items()}
    )
    periodos: dict = field(default_factory=lambda: dict(PERIODOS_ACADEMICOS))
    anio_min: int = PERIODO_ANIO_MINIMO
    anio_max: int = PERIODO_ANIO_MAXIMO
    leer_catalogos_desde_hoja: bool = LEER_CATALOGOS_DESDE_HOJA
    registrar_validos: bool = REGISTRAR_VALIDOS_EN_REPORTE
    # --- Modo "convertir a plantilla" (ver convertir_plantilla.py) ---
    mapeo: list = field(default_factory=list)            # [{"destino","origen","regla"}]
    periodos_permitidos: list = field(default_factory=list)  # ["202710", ...]
    periodo_sufijos: dict = field(default_factory=dict)  # {"1": "10", "2": "20"}

    @property
    def columnas_catalogo(self):
        return [self.col_nivel, self.col_documento, self.col_programa]


def cargar_configuracion_json(ruta):
    """Crea una Configuracion a partir de un JSON (ver config_prueba.json)."""
    datos = json.loads(Path(ruta).read_text(encoding="utf-8"))
    cfg = Configuracion()
    for clave, valor in datos.items():
        if clave.startswith("_"):
            continue
        if not hasattr(cfg, clave):
            raise ValueError(f"Clave de configuración desconocida: {clave}")
        setattr(cfg, clave, valor)
    return cfg


# ============================================================
# 2. FUNCIONES AUXILIARES
# ============================================================

def normalizar_comparacion(valor):
    """
    Clave de comparación: minúsculas, sin tildes, sin "_", espacios
    colapsados. La ñ se conserva (no es una tilde: "año" != "ano").
    """
    if valor is None:
        return ""
    texto = str(valor).replace("\u00a0", " ").strip().lower()
    texto = texto.replace("_", " ").replace("ñ", "\0")
    texto = re.sub(r"[,;:.]", " ", texto)  # puntuación sin significado
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.replace("\0", "ñ")
    return re.sub(r"\s+", " ", texto).strip()


normalizar_encabezado = normalizar_comparacion


def es_vacio(valor):
    return valor is None or (isinstance(valor, str) and valor.strip() == "")


def convertir_title_case(valor):
    """
    "_" -> espacio, espacios colapsados y Title Case por palabra.
    No agrega tildes. CONSULTA__EXTERNA -> Consulta Externa
    """
    texto = str(valor).replace("_", " ")
    texto = re.sub(r"\s+", " ", texto).strip()
    return re.sub(
        r"(?<!\w)(\w)(\w*)",
        lambda m: m.group(1).upper() + m.group(2).lower(),
        texto,
    )


def buscar_columnas(hoja, nombre_columna):
    """Devuelve todas las posiciones cuyo encabezado coincide."""
    objetivo = normalizar_encabezado(nombre_columna)
    return [
        celda.column
        for celda in hoja[1]
        if normalizar_encabezado(celda.value) == objetivo
    ]


class Auditoria:
    """Acumula filas del reporte y contadores del resumen."""

    def __init__(self, registrar_validos):
        self.filas = []
        self.contador = Counter()
        self.registrar_validos = registrar_validos
        self.sin_catalogo = Counter()

    def contar(self, hoja, columna, estado):
        self.contador[(hoja, columna, estado)] += 1

    def registrar(self, hoja, fila, columna, original, final, resultado, motivo=""):
        self.contar(hoja, columna, resultado)
        if resultado == "VALIDO" and not self.registrar_validos:
            return
        self.filas.append({
            "Hoja": hoja,
            "Fila": fila,
            "Columna": columna,
            "Valor original": "" if original is None else str(original),
            "Valor final": "" if final is None else str(final),
            "Resultado": resultado,
            "Motivo": motivo,
        })


# ============================================================
# 3. VALIDACIÓN CONTRA CATÁLOGOS
# ============================================================

def validar_valor(valor, catalogo, equivalencias=None):
    """
    Devuelve (valor_final, estado, motivo).

    Estados: VALIDO, CORREGIDO, INVALIDO, REQUIERE_REVISION,
    SIN_CATALOGO, VACIO.

    Solo corrige si hay UNA coincidencia inequívoca tras normalizar
    mayúsculas, tildes y espacios, o si la equivalencia está
    confirmada en la configuración. Nunca usa similitud aproximada.
    """
    equivalencias = equivalencias or {}

    if es_vacio(valor):
        return valor, "VACIO", ""

    clave = normalizar_comparacion(valor)
    equiv_norm = {normalizar_comparacion(k): v for k, v in equivalencias.items()}

    # 1) Equivalencia confirmada (abreviaturas, alias)
    if clave in equiv_norm:
        oficial = equiv_norm[clave]
        if catalogo and oficial not in catalogo:
            return (
                valor,
                "REQUIERE_REVISION",
                f"La equivalencia configurada apunta a '{oficial}', que no está en el catálogo.",
            )
        if str(valor) == oficial:
            return valor, "VALIDO", ""
        return oficial, "CORREGIDO", "Equivalencia confirmada en la configuración."

    # 2) Sin catálogo no se corrige nada
    if not catalogo:
        return (
            valor,
            "SIN_CATALOGO",
            "La lista de valores válidos aún no está configurada.",
        )

    # 3) Coincidencia normalizada con el catálogo
    coincidencias = list(dict.fromkeys(
        op for op in catalogo if normalizar_comparacion(op) == clave
    ))
    if len(coincidencias) == 1:
        oficial = coincidencias[0]
        if str(valor) == oficial:
            return valor, "VALIDO", ""
        return oficial, "CORREGIDO", "Se ajustó a la escritura oficial del catálogo."
    if len(coincidencias) > 1:
        return (
            valor,
            "REQUIERE_REVISION",
            "Varias opciones del catálogo son equivalentes; no se puede elegir una.",
        )
    return valor, "INVALIDO", "El valor no coincide con ninguna opción del catálogo."


# ============================================================
# 4. TELÉFONOS
# ============================================================

def normalizar_telefono(valor):
    """
    Devuelve (telefono_final, codigo_pais, modificado, motivo_revision).

    Solo actúa si el texto empieza explícitamente por "+57" seguido de
    un número nacional. No toca teléfonos numéricos ni agrega +57.
    """
    if es_vacio(valor) or not isinstance(valor, str):
        return valor, None, False, ""

    texto = valor.replace("\u00a0", " ").strip()
    if not texto.startswith("+57"):
        return valor, None, False, ""

    nacional = texto[3:].strip()
    if not nacional:
        return valor, None, False, (
            "El teléfono contiene solo +57, sin número nacional."
        )
    if not nacional[0].isdigit():
        return valor, None, False, (
            "Formato inesperado después de +57; se conserva el original."
        )
    return nacional, "+57", True, ""


# ============================================================
# 5. PERIODOS
# ============================================================

def normalizar_periodo(valor, periodos, anio_min=2000, anio_max=2100):
    """
    202710 -> "202710 - Primer Semestre 2027".

    Devuelve (valor_final, estado, motivo). Estados: VACIO, VALIDO,
    CORREGIDO, REQUIERE_REVISION. Los códigos no configurados o los
    formatos desconocidos se conservan sin inventar descripciones.
    """
    if es_vacio(valor):
        return valor, "VACIO", ""

    if isinstance(valor, bool):
        return valor, "REQUIERE_REVISION", "Tipo de dato inesperado."
    if isinstance(valor, float) and valor.is_integer():
        texto = str(int(valor))
    elif isinstance(valor, int):
        texto = str(valor)
    else:
        texto = str(valor).replace("\u00a0", " ").strip()

    def _evaluar_codigo(anio_txt, cod):
        anio = int(anio_txt)
        if not (anio_min <= anio <= anio_max):
            return None, f"Año {anio} fuera del rango permitido ({anio_min}-{anio_max})."
        desc = periodos.get(cod)
        if desc is None:
            return None, f"No hay descripción configurada para el código {cod}."
        return f"{anio_txt}{cod} - {desc} {anio_txt}", ""

    m = re.fullmatch(r"(\d{4})(\d{2})", texto)
    if m:
        canonico, motivo = _evaluar_codigo(m.group(1), m.group(2))
        if canonico is None:
            return valor, "REQUIERE_REVISION", motivo
        return canonico, "CORREGIDO", "Código convertido a código + descripción + año."

    m = re.fullmatch(r"(\d{4})(\d{2})\s*-\s*(.+?)\s+(\d{4})", texto)
    if m:
        anio_txt, cod, desc, anio_desc = m.groups()
        canonico, motivo = _evaluar_codigo(anio_txt, cod)
        if canonico is None:
            return valor, "REQUIERE_REVISION", motivo
        if anio_desc != anio_txt or normalizar_comparacion(desc) != normalizar_comparacion(
            periodos[cod]
        ):
            return valor, "REQUIERE_REVISION", (
                "La descripción o el año no coinciden con el código; se conserva el original."
            )
        if texto == canonico:
            return valor, "VALIDO", ""
        return canonico, "CORREGIDO", "Formato de periodo normalizado."

    return valor, "REQUIERE_REVISION", "Formato no reconocido; se conserva el valor original."


# ============================================================
# 6. LECTURA OPCIONAL DE CATÁLOGOS DESDE "Reglas de validación"
# ============================================================

def leer_catalogos_desde_hoja(libro, cfg):
    """
    Lee catálogos de la hoja "Reglas de validación" SOLO para las
    columnas cuyo encabezado (fila 1) coincida con una columna
    configurada. Devuelve los nombres de columnas cargadas.
    """
    objetivo = normalizar_encabezado(HOJA_REGLAS)
    hoja = next(
        (h for h in libro.worksheets if normalizar_encabezado(h.title) == objetivo),
        None,
    )
    if hoja is None:
        return []

    cargadas = []
    for celda in hoja[1]:
        for nombre in cfg.columnas_catalogo:
            if normalizar_encabezado(celda.value) == normalizar_encabezado(nombre):
                valores = []
                for fila in range(2, hoja.max_row + 1):
                    v = hoja.cell(row=fila, column=celda.column).value
                    if not es_vacio(v):
                        valores.append(str(v).strip())
                if valores:
                    cfg.valores_validos[nombre] = list(dict.fromkeys(valores))
                    cargadas.append(nombre)
    return cargadas


# ============================================================
# 7. PROCESAMIENTO
# ============================================================

def _es_celda_editable(celda):
    return not isinstance(celda, MergedCell) and celda.data_type != "f"


def procesar_excel(entrada=ARCHIVO_ENTRADA, salida=ARCHIVO_SALIDA, cfg=None):
    cfg = cfg or Configuracion()
    entrada, salida = Path(entrada), Path(salida)

    if not entrada.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {entrada.resolve()}")

    es_xlsm = entrada.suffix.lower() == ".xlsm"
    if es_xlsm and salida.suffix.lower() != ".xlsm":
        salida = salida.with_suffix(".xlsm")  # evita destruir las macros
    if entrada.resolve() == salida.resolve():
        raise ValueError("El archivo de salida debe ser distinto del original.")

    # Se cargan las fórmulas (no sus resultados) para conservarlas.
    libro = load_workbook(entrada, keep_vba=es_xlsm)

    auditoria = Auditoria(cfg.registrar_validos)
    reg = auditoria.registrar
    avisos = []

    if cfg.leer_catalogos_desde_hoja:
        cargadas = leer_catalogos_desde_hoja(libro, cfg)
        avisos.append(f"Catálogos leídos de '{HOJA_REGLAS}': {cargadas or 'ninguno'}")

    hojas_omitidas = {
        normalizar_encabezado(n) for n in (HOJA_REPORTE, HOJA_RESUMEN, HOJA_REGLAS)
    }
    hojas_procesadas = 0

    for hoja in libro.worksheets:
        if normalizar_encabezado(hoja.title) in hojas_omitidas:
            continue
        hojas_procesadas += 1

        if getattr(hoja, "_charts", None) or getattr(hoja, "_images", None):
            avisos.append(
                f"La hoja '{hoja.title}' tiene gráficos o imágenes que openpyxl "
                "podría no conservar al guardar."
            )

        # ---- Ubicar columnas (se rechazan encabezados duplicados) ----
        def ubicar(nombre):
            posiciones = buscar_columnas(hoja, nombre)
            if len(posiciones) > 1:
                reg(hoja.title, "", nombre, "", "", "REQUIERE_REVISION",
                    f"Encabezado duplicado en las columnas {posiciones}; no se procesa.")
                return None
            return posiciones[0] if posiciones else None

        col_catalogo = {n: ubicar(n) for n in cfg.columnas_catalogo}
        col_texto = ubicar(cfg.col_texto)
        col_tel = ubicar(cfg.col_telefono)
        col_codigo = ubicar(cfg.col_codigo_pais)
        col_periodo = ubicar(cfg.col_periodo)

        esperadas = {**col_catalogo, cfg.col_texto: col_texto,
                     cfg.col_telefono: col_tel, cfg.col_periodo: col_periodo}
        for nombre, pos in esperadas.items():
            if pos is None and not buscar_columnas(hoja, nombre):
                reg(hoja.title, "", nombre, "", "", "COLUMNA_NO_ENCONTRADA",
                    "Verifica el encabezado configurado.")

        def asegurar_columna_codigo():
            """Crea Codigo_Pais solo cuando hace falta."""
            nonlocal col_codigo
            if col_codigo is not None:
                return col_codigo
            col_codigo = hoja.max_column + 1
            enc = hoja.cell(row=1, column=col_codigo, value=cfg.col_codigo_pais)
            ref = hoja.cell(row=1, column=col_tel)
            if ref.has_style:
                enc._style = copy(ref._style)
            return col_codigo

        # ---- Filas ----
        for fila in range(2, hoja.max_row + 1):

            # 1. Catálogos
            for nombre, pos in col_catalogo.items():
                if pos is None:
                    continue
                celda = hoja.cell(row=fila, column=pos)
                original = celda.value
                if es_vacio(original):
                    continue
                if not _es_celda_editable(celda):
                    reg(hoja.title, fila, nombre, original, original,
                        "REQUIERE_REVISION", "Celda con fórmula o combinada; no se modifica.")
                    continue
                final, estado, motivo = validar_valor(
                    original,
                    cfg.valores_validos.get(nombre, []),
                    cfg.equivalencias.get(nombre, {}),
                )
                if estado == "CORREGIDO":
                    celda.value = final
                if estado == "SIN_CATALOGO":
                    auditoria.sin_catalogo[(hoja.title, nombre)] += 1
                    auditoria.contar(hoja.title, nombre, estado)
                else:
                    reg(hoja.title, fila, nombre, original, final, estado, motivo)

            # 2. categoria_servicio
            if col_texto is not None:
                celda = hoja.cell(row=fila, column=col_texto)
                original = celda.value
                if isinstance(original, str) and not es_vacio(original):
                    if not _es_celda_editable(celda):
                        reg(hoja.title, fila, cfg.col_texto, original, original,
                            "REQUIERE_REVISION", "Celda con fórmula o combinada.")
                    else:
                        final = convertir_title_case(original)
                        if final != original:
                            celda.value = final
                            reg(hoja.title, fila, cfg.col_texto, original, final,
                                "CORREGIDO",
                                "Guiones bajos reemplazados, espacios depurados y Title Case.")
                        else:
                            reg(hoja.title, fila, cfg.col_texto, original, final, "VALIDO")

            # 3. Teléfono / indicativo
            if col_tel is not None:
                c_tel = hoja.cell(row=fila, column=col_tel)
                original = c_tel.value
                if _es_celda_editable(c_tel):
                    final, codigo, modificado, motivo = normalizar_telefono(original)
                    if modificado:
                        existente = None
                        if col_codigo is not None:
                            existente = hoja.cell(row=fila, column=col_codigo).value
                        c_cod = hoja.cell(row=fila, column=col_codigo) if col_codigo else None
                        if c_cod is not None and not _es_celda_editable(c_cod):
                            reg(hoja.title, fila, cfg.col_codigo_pais, existente, existente,
                                "REQUIERE_REVISION", "Celda de código con fórmula o combinada.")
                        elif not es_vacio(existente) and str(existente).strip() != "+57":
                            reg(hoja.title, fila, cfg.col_codigo_pais, existente, existente,
                                "REQUIERE_REVISION",
                                f"Ya existe un código distinto de +57 (teléfono: {original}); "
                                "no se modificó ninguna de las dos celdas.")
                        else:
                            c_cod = hoja.cell(row=fila, column=asegurar_columna_codigo())
                            c_tel.value = final
                            c_tel.data_type = "s"
                            c_cod.value = "+57"
                            c_cod.data_type = "s"
                            reg(hoja.title, fila, cfg.col_telefono, original, final,
                                "CORREGIDO", "Indicativo +57 separado en otra columna.")
                            reg(hoja.title, fila, cfg.col_codigo_pais, existente, "+57",
                                "CORREGIDO", "Código de país asignado desde el teléfono.")
                    elif motivo:
                        reg(hoja.title, fila, cfg.col_telefono, original, original,
                            "REQUIERE_REVISION", motivo)

            # 4. Periodo
            if col_periodo is not None:
                celda = hoja.cell(row=fila, column=col_periodo)
                original = celda.value
                if es_vacio(original):
                    continue
                if not _es_celda_editable(celda):
                    reg(hoja.title, fila, cfg.col_periodo, original, original,
                        "REQUIERE_REVISION", "Celda con fórmula o combinada.")
                    continue
                final, estado, motivo = normalizar_periodo(
                    original, cfg.periodos, cfg.anio_min, cfg.anio_max
                )
                if estado == "CORREGIDO":
                    celda.value = final
                    celda.data_type = "s"
                reg(hoja.title, fila, cfg.col_periodo, original, final, estado, motivo)

    # ---- Una fila de resumen por hoja/columna sin catálogo ----
    for (nombre_hoja, nombre_col), n in sorted(auditoria.sin_catalogo.items()):
        auditoria.filas.append({
            "Hoja": nombre_hoja, "Fila": "", "Columna": nombre_col,
            "Valor original": f"{n} celda(s) con valor", "Valor final": "",
            "Resultado": "SIN_CATALOGO",
            "Motivo": "Catálogo oficial pendiente: los valores se conservaron sin cambios.",
        })

    _escribir_reportes(libro, auditoria)
    libro.save(salida)

    resumen = Counter()
    for (_, _, estado), n in auditoria.contador.items():
        resumen[estado] += n
    return {
        "salida": salida,
        "hojas_procesadas": hojas_procesadas,
        "reporte": auditoria.filas,
        "resumen": dict(resumen),
        "avisos": avisos,
    }


def _escribir_reportes(libro, auditoria):
    for nombre in (HOJA_REPORTE, HOJA_RESUMEN):
        if nombre in libro.sheetnames:
            del libro[nombre]

    # --- Reporte detallado ---
    enc = ["Hoja", "Fila", "Columna", "Valor original", "Valor final", "Resultado", "Motivo"]
    hoja = libro.create_sheet(HOJA_REPORTE)
    hoja.append(enc)
    for r in auditoria.filas:
        hoja.append([r[c] for c in enc])
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = "s"  # evita que "=..." se interprete como fórmula
    _formato_tabla(hoja)

    # --- Resumen por hoja / columna / estado ---
    res = libro.create_sheet(HOJA_RESUMEN)
    res.append(["Hoja", "Columna", "Resultado", "Cantidad"])
    for (h, c, e), n in sorted(auditoria.contador.items()):
        res.append([h, c, e, n])
    _formato_tabla(res)


def _formato_tabla(hoja):
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for columna in hoja.columns:
        largo = max(
            (len(str(c.value)) for c in columna if c.value is not None), default=10
        )
        hoja.column_dimensions[columna[0].column_letter].width = min(largo + 2, 60)


# ============================================================
# 8. PUNTO DE ENTRADA
# ============================================================

def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    p.add_argument("--entrada", default=str(ARCHIVO_ENTRADA))
    p.add_argument("--salida", default=str(ARCHIVO_SALIDA))
    p.add_argument("--config", help="JSON con catálogos/equivalencias (opcional)")
    args = p.parse_args(argv)

    cfg = cargar_configuracion_json(args.config) if args.config else Configuracion()
    r = procesar_excel(args.entrada, args.salida, cfg)

    print("\n" + "=" * 55)
    print("PROCESAMIENTO COMPLETADO")
    print("=" * 55)
    print(f"Archivo original : {args.entrada}")
    print(f"Archivo generado : {r['salida']}")
    print(f"Hojas procesadas : {r['hojas_procesadas']}")
    for estado, n in sorted(r["resumen"].items()):
        print(f"  {estado:<22}{n}")
    print(f"Filas en el reporte: {len(r['reporte'])}")
    for aviso in r["avisos"]:
        print(f"AVISO: {aviso}")
    print("=" * 55)


if __name__ == "__main__":
    main()
