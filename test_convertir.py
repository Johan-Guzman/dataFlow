import hashlib
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

import automatizacion_excel as ae
import convertir_plantilla as cp

AQUI = Path(__file__).parent
GUIA = AQUI / "guia_anonimizada.xlsx"
PLANTILLA = AQUI / "plantilla_importacion_Aspirante.xlsx"
CONFIG = AQUI / "config_aspirantes.json"


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def correr(tmp_path, entrada=GUIA, cfg=None):
    cfg = cfg or ae.cargar_configuracion_json(CONFIG)
    salida, rep = tmp_path / "imp.xlsx", tmp_path / "rep.xlsx"
    r = cp.convertir_a_plantilla(entrada, PLANTILLA, salida, rep, cfg)
    ws = load_workbook(salida).worksheets[0]
    cab = [c.value for c in ws[1]]
    filas = [{h: ws.cell(row=i, column=j + 1).value for j, h in enumerate(cab)} for i in range(2, ws.max_row + 1)]
    return r, ws, filas


def guia_modificada(tmp_path, cambios):
    """Copia la guía y cambia celdas: {(fila, 'encabezado'): valor}."""
    ruta = tmp_path / "sucio.xlsx"
    shutil.copy(GUIA, ruta)
    from openpyxl import load_workbook as lw
    wb = lw(ruta)
    ws = wb.worksheets[0]
    cab = {c.value: c.column for c in ws[1]}
    for (fila, h), v in cambios.items():
        ws.cell(row=fila, column=cab[h]).value = v
    wb.save(ruta)
    return ruta


def test_estructura_y_filas(tmp_path):
    r, ws, filas = correr(tmp_path)
    plantilla = load_workbook(PLANTILLA).worksheets[0]
    assert [c.value for c in ws[1]] == [c.value for c in plantilla[1]]
    assert load_workbook(tmp_path / "imp.xlsx").sheetnames == ["Importacion"]  # solo la hoja de datos
    assert len(filas) == 26 == r["filas_origen"] == r["filas_salida"]
    assert all(f["Nombre"] != "Camilo" for f in filas)  # se quitó el ejemplo de la plantilla


def test_catalogos_y_normalizacion(tmp_path):
    _, _, filas = correr(tmp_path)
    assert filas[0]["Tipo de documento"] == "CC - Cédula de ciudadanía"
    assert filas[0]["Nivel educativo"] == "Profesional"
    assert filas[0]["Programa de interés"] == "Especialización en Administración en Salud"  # guion bajo final
    assert filas[1]["Programa de interés"] == "MBA Icesi"
    programas = {f["Programa de interés"] for f in filas}
    assert "Maestría en Discapacidad Salud y Derechos Humanos" in programas  # coma ignorada
    assert "Maestría en Gerencia en Salud" in programas
    assert "Programa de Educación Continua" in programas
    assert "Otro" in programas


def test_telefono_y_documento(tmp_path):
    _, _, filas = correr(tmp_path)
    assert filas[0]["Código de país"] == "+57"
    assert filas[0]["Número de teléfono"] == "0505050505"  # conserva el cero inicial
    assert isinstance(filas[0]["Número de teléfono"], str)
    assert filas[0]["Número de documento"] == "10100000000"
    assert isinstance(filas[0]["Número de documento"], str)


def test_no_inventa_valores_sin_equivalencia(tmp_path):
    r, ws, filas = correr(tmp_path)
    niveles = [f["Nivel educativo"] for f in filas]
    assert niveles.count("especialización") == 15 and niveles.count("maestría") == 3
    assert "Especialista" not in niveles and "Magíster" not in niveles
    prog = [f["Programa de interés"] for f in filas]
    assert "maestría_en_ciencia_de_datos" in prog  # NO se convierte en "Ciencias de Datos"
    assert "pe_-_permiso_especial_de_permanencia" in [f["Tipo de documento"] for f in filas]
    inval = [x for x in r["reporte"] if x["Resultado"] == "INVALIDO"]
    assert len(inval) == 30
    assert r["resumen"]["INVALIDO"] == 30


def test_equivalencia_confirmada_se_aplica(tmp_path):
    cfg = ae.cargar_configuracion_json(CONFIG)
    cfg.equivalencias["nivel educativo"] = {"especialización": "Especialista", "maestría": "Magíster"}
    _, _, filas = correr(tmp_path, cfg=cfg)
    niveles = {f["Nivel educativo"] for f in filas}
    assert niveles == {"Profesional", "Especialista", "Magíster"}


def test_periodo(tmp_path):
    r, _, filas = correr(tmp_path)
    assert all(f["¿En qué semestre te interesaría iniciar el Programa? "] == "2026-1" for f in filas)
    motivos = {x["Motivo"] for x in r["reporte"] if x["Resultado"] == "REQUIERE_REVISION"}
    assert any("202610" in m and "permitidos" in m for m in motivos)


def test_periodo_formato_alterno_valido(tmp_path):
    h = "Periodo de interes"
    ruta = guia_modificada(tmp_path, {(2, h): "2027-1", (3, h): "2027-2", (4, h): "202810",
                                      (5, h): "periodo pendiente", (6, h): "2029-1"})
    _, _, filas = correr(tmp_path, ruta)
    col = "¿En qué semestre te interesaría iniciar el Programa? "
    assert filas[0][col] == "202710 - Primer Semestre 2027"
    assert filas[1][col] == "202720 - Segundo Semestre 2027"
    assert filas[2][col] == "202810 - Primer Semestre 2028"
    assert filas[3][col] == "periodo pendiente"
    assert filas[4][col] == "2029-1"  # 202910 no está permitido


def test_no_usa_columna_de_evento_como_programa(tmp_path):
    _, _, filas = correr(tmp_path)
    assert all("Icesi Conecta" not in str(v) for f in filas for v in f.values())


def test_vacios_y_telefonos_especiales(tmp_path):
    ruta = guia_modificada(tmp_path, {(2, "city"): None, (3, "phone_number"): "3009876543",
                                      (4, "phone_number"): "+57", (5, "phone_number"): "+1 5551234567",
                                      (6, "email"): "correo-sin-arroba"})
    r, ws, filas = correr(tmp_path, ruta)
    assert filas[0]["Ciudad de residencia"] is None
    assert (filas[1]["Código de país"], filas[1]["Número de teléfono"]) == (None, "3009876543")
    assert (filas[2]["Código de país"], filas[2]["Número de teléfono"]) == (None, "+57")
    assert (filas[3]["Código de país"], filas[3]["Número de teléfono"]) == (None, "+1 5551234567")
    revis = {(x["Fila"], x["Columna"]) for x in r["reporte"] if x["Resultado"] == "REQUIERE_REVISION"}
    assert (4, "Número de teléfono") in revis and (5, "Número de teléfono") in revis
    assert (6, "Correo electrónico") in revis
    assert ws["C6"].fill.fgColor.rgb.endswith("FFF1D6")  # celda resaltada


def test_filas_vacias_se_omiten_y_orden_se_conserva(tmp_path):
    ruta = tmp_path / "s.xlsx"
    shutil.copy(GUIA, ruta)
    wb = load_workbook(ruta); ws = wb.worksheets[0]
    for c in ws[5]:  # fila vacía en medio (los correos traen hipervínculo: también se quita)
        c.hyperlink = None
        c.value = None
    wb.save(ruta)
    r, _, filas = correr(tmp_path, ruta)
    assert len(filas) == 25 == r["filas_origen"]  # la fila vacía se omite, no queda hueco


def test_origen_y_plantilla_no_se_modifican(tmp_path):
    a, b = sha(GUIA), sha(PLANTILLA)
    correr(tmp_path)
    assert (sha(GUIA), sha(PLANTILLA)) == (a, b)


def test_reporte_en_archivo_aparte(tmp_path):
    correr(tmp_path)
    rep = load_workbook(tmp_path / "rep.xlsx")
    assert set(rep.sheetnames) == {"Reporte_validacion", "Resumen_validacion"}
    assert [c.value for c in rep["Reporte_validacion"][1]] == [
        "Hoja", "Fila", "Columna", "Valor original", "Valor final", "Resultado", "Motivo"]


def test_columna_origen_inexistente(tmp_path):
    cfg = ae.cargar_configuracion_json(CONFIG)
    cfg.mapeo[0]["origen"] = "no_existe"
    r, _, filas = correr(tmp_path, cfg=cfg)
    assert filas[0]["Nombre"] is None
    assert any(x["Resultado"] == "COLUMNA_NO_ENCONTRADA" for x in r["reporte"])
