"""Lógica compartida por las interfaces (sin dependencias de pantalla)."""

import json
import tempfile
from pathlib import Path

import automatizacion_excel as ae
import convertir_plantilla as cp

COLUMNAS_CATALOGO = [
    ("nivel", "Nivel educativo"),
    ("documento", "Tipo de documento"),
    ("programa", "Programa de interés"),
]
ENCABEZADOS = [
    ("col_nivel", "Nivel educativo"),
    ("col_documento", "Tipo de documento"),
    ("col_programa", "Programa de interés"),
    ("col_texto", "Texto a normalizar (Title Case)"),
    ("col_telefono", "Teléfono"),
    ("col_codigo_pais", "Código de país"),
    ("col_periodo", "Periodo"),
]


def lineas(texto):
    return [l.strip() for l in texto.splitlines() if l.strip()]


def parsear_pares(texto, etiqueta):
    """'clave = valor' por línea -> dict. Rechaza líneas mal formadas y conflictos."""
    pares, vistos = {}, {}
    for linea in lineas(texto):
        if "=" not in linea:
            raise ValueError(f"{etiqueta}: la línea «{linea}» debe tener formato clave = valor")
        k, v = (x.strip() for x in linea.split("=", 1))
        if not k or not v:
            raise ValueError(f"{etiqueta}: la línea «{linea}» tiene la clave o el valor vacío")
        nk = ae.normalizar_comparacion(k)
        if nk in vistos and pares[vistos[nk]] != v:
            raise ValueError(
                f"{etiqueta}: «{vistos[nk]}» y «{k}» son equivalentes pero apuntan a valores distintos"
            )
        vistos[nk] = k
        pares[k] = v
    return pares


def pares_a_texto(d):
    return "\n".join(f"{k} = {v}" for k, v in d.items())


def estado_desde_cfg(cfg):
    """Configuracion -> dict plano con las claves que usa el formulario."""
    e = {clave: getattr(cfg, clave) for clave, _ in ENCABEZADOS}
    for clave, _ in COLUMNAS_CATALOGO:
        nombre = getattr(cfg, f"col_{clave}")
        e[f"cat_{clave}"] = "\n".join(cfg.valores_validos.get(nombre, []))
        e[f"eq_{clave}"] = pares_a_texto(cfg.equivalencias.get(nombre, {}))
    e.update(
        periodos=pares_a_texto(cfg.periodos),
        anio_min=int(cfg.anio_min),
        anio_max=int(cfg.anio_max),
        leer_hoja=bool(cfg.leer_catalogos_desde_hoja),
        reg_validos=bool(cfg.registrar_validos),
        per_permitidos="\n".join(cfg.periodos_permitidos),
        per_sufijos=pares_a_texto(cfg.periodo_sufijos),
        mapeo=[dict(m) for m in cfg.mapeo],
    )
    for i, m in enumerate(cfg.mapeo):
        e[f"map_{i}"] = m.get("origen", "")
    return e


def construir_cfg(e):
    """Dict/estado del formulario -> Configuracion validada (ValueError si hay errores)."""
    cfg = ae.Configuracion()
    nombres = []
    for clave, etiqueta in ENCABEZADOS:
        # En modo plantilla estos campos no se dibujan y Streamlit descarta su estado
        valor = e.get(clave)
        valor = str(getattr(cfg, clave) if valor is None else valor).strip()
        if not valor:
            raise ValueError(f"El encabezado «{etiqueta}» no puede estar vacío.")
        setattr(cfg, clave, valor)
        nombres.append(ae.normalizar_encabezado(valor))
    if len(set(nombres)) != len(nombres):
        raise ValueError("Hay dos columnas con el mismo encabezado.")

    cfg.valores_validos, cfg.equivalencias = {}, {}
    for clave, etiqueta in COLUMNAS_CATALOGO:
        nombre = getattr(cfg, f"col_{clave}")
        cfg.valores_validos[nombre] = lineas(e[f"cat_{clave}"])
        cfg.equivalencias[nombre] = parsear_pares(e[f"eq_{clave}"], f"Equivalencias de {etiqueta}")

    cfg.periodos = parsear_pares(e["periodos"], "Periodos")
    for cod in cfg.periodos:
        if not (cod.isdigit() and len(cod) == 2):
            raise ValueError(f"Periodos: el código «{cod}» debe tener exactamente 2 dígitos.")
    cfg.anio_min, cfg.anio_max = int(e["anio_min"]), int(e["anio_max"])
    if cfg.anio_min > cfg.anio_max:
        raise ValueError("El año mínimo no puede ser mayor que el máximo.")
    cfg.leer_catalogos_desde_hoja = bool(e["leer_hoja"])
    cfg.registrar_validos = bool(e["reg_validos"])

    # Modo "convertir a plantilla" (opcionales; vacíos en el modo normal)
    cfg.periodos_permitidos = lineas(str(e.get("per_permitidos", "")))
    for cod in cfg.periodos_permitidos:
        if not (cod.isdigit() and len(cod) == 6):
            raise ValueError(f"Periodos permitidos: «{cod}» debe tener 6 dígitos (ej. 202710).")
    cfg.periodo_sufijos = parsear_pares(str(e.get("per_sufijos", "")), "Sufijos de periodo")
    cfg.mapeo = [
        {**m, "origen": str(e.get(f"map_{i}", m.get("origen", ""))).strip()}
        for i, m in enumerate(e.get("mapeo", []))
    ]
    return cfg


def cfg_a_json(cfg):
    from dataclasses import asdict
    return json.dumps(asdict(cfg), ensure_ascii=False, indent=2)


def cfg_desde_json_bytes(contenido):
    with tempfile.TemporaryDirectory() as d:
        ruta = Path(d) / "config.json"
        ruta.write_bytes(contenido)
        return ae.cargar_configuracion_json(ruta)


def procesar_bytes(contenido, nombre, cfg):
    """Procesa un Excel en memoria. Devuelve el resultado + bytes del archivo corregido."""
    entrada_nombre = Path(nombre).name
    sufijo = Path(entrada_nombre).suffix.lower()
    if sufijo not in (".xlsx", ".xlsm"):
        raise ValueError("Solo se aceptan archivos .xlsx o .xlsm.")
    with tempfile.TemporaryDirectory() as d:
        entrada = Path(d) / entrada_nombre
        entrada.write_bytes(contenido)
        salida = Path(d) / f"{entrada.stem}_corregido{sufijo}"
        r = ae.procesar_excel(entrada, salida, cfg)
        r["bytes"] = Path(r["salida"]).read_bytes()
        r["nombre_salida"] = Path(r["salida"]).name
    del r["salida"]
    return r


def convertir_bytes(contenido, nombre, plantilla, cfg):
    """Convierte un Excel en memoria a la plantilla. Devuelve archivos listo + reporte (bytes)."""
    entrada_nombre = Path(nombre).name
    if Path(entrada_nombre).suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("Solo se aceptan archivos .xlsx o .xlsm.")
    with tempfile.TemporaryDirectory() as d:
        entrada = Path(d) / entrada_nombre
        entrada.write_bytes(contenido)
        base = Path(entrada_nombre).stem
        salida, reporte = Path(d) / f"{base}_importacion.xlsx", Path(d) / f"{base}_reporte.xlsx"
        r = cp.convertir_a_plantilla(entrada, plantilla, salida, reporte, cfg)
        r["bytes"], r["nombre_salida"] = salida.read_bytes(), salida.name
        r["bytes_reporte"], r["nombre_reporte"] = reporte.read_bytes(), reporte.name
        # Añadir vista previa del archivo convertido
        try:
            from openpyxl import load_workbook
            import pandas as pd
            wb_prev = load_workbook(salida)
            ws_prev = wb_prev.worksheets[0]
            cab = [c.value for c in ws_prev[1]]
            datos = []
            for i in range(2, min(ws_prev.max_row + 1, 102)):  # limitar vista previa
                fila_d = {}
                for j, h in enumerate(cab):
                    fila_d[h] = ws_prev.cell(row=i, column=j + 1).value
                datos.append(fila_d)
            r["preview_df"] = datos
            r["preview_headers"] = cab
            r["archivo_origen"] = entrada_nombre
            r["registros_detectados"] = r.get("filas_origen", len(datos))
            r["registros_exportar"] = r.get("filas_salida", len(datos))
            # Mapeo origen->destino
            r["mapeo_origen_destino"] = [{"origen": m.get("origen"), "destino": m.get("destino"), "regla": m.get("regla")} for m in getattr(cfg, "mapeo", [])]
        except Exception:
            pass
    del r["salida"], r["salida_reporte"]
    return r
