"""Genera un Excel de prueba con datos ficticios."""
from openpyxl import Workbook

ENCABEZADOS = [
    "id", "nombre", "nivel educativo", "tipo de documento", "programa de interés",
    "categoria_servicio", "telefono", "Codigo_Pais", "periodo",
    "observaciones", "cantidad", "total",
]

FILAS = [
    # id, nombre, nivel, doc, programa, categoria, tel, cod, periodo, obs, cant
    (1, "Ana", "bachillerato", "CC", "Ingenieria de Sistemas", "CONSULTA__EXTERNA", "+573001234567", None, "202710", "TEXTO_LIBRE sin Tocar", 3),
    (2, "Luis", "TECNICO", "TI", "Ingeniería de Sistemas", "MEDICINA_GENERAL", "+57 312 555 0182", None, "202720", "otra_nota", 5),
    (3, "Eva", "Tecnólogo", "CE", "Programa no catalogado", "ATENCION_INTEGRAL_EN_SALUD", "3009876543", None, "202830", "x", 1),
    (4, "Juan", "Profesional", "Cédula de ciudadanía", "Derecho", "SERVICIOS_DE_URGENCIAS", "+573155551234", None, "202810", "y", 2),
    (5, "Sara", "Doctorado", "Pasaporte", "Medicina", "ODONTOLOGIA", "+57", None, "periodo pendiente", "z", 4),
    (6, "Leo", "bachillerato ", "RC", "Ingenieria de Sistemas", "CONSULTA_EXTERNA", "+573001234567", "+1", "202710 - Primer Semestre 2027", "w", 6),
    (7, "Vacío", None, None, None, None, None, None, None, None, None),
    (8, "Mia", "Profesional", "TI", "Derecho", "APOYO_PSICOSOCIAL", 3105551234, None, 202720, "n", 7),
    (9, "Teo", "Técnico", "CC", "Medicina", "  apoyo__psicosocial ", "0123456789", "+57", "202710 - Segundo Semestre 2027", "k", 8),
]


def crear_excel_prueba(ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(ENCABEZADOS)
    for i, fila in enumerate(FILAS, start=2):
        ws.append(list(fila) + [f"=K{i}*2"])
    # formato de texto para teléfonos (conserva ceros iniciales)
    for i in range(2, len(FILAS) + 2):
        ws.cell(row=i, column=7).number_format = "@"
    ws["K2"].number_format = "0.00"
    # segunda hoja con otra columna que NO debe tocarse
    ws2 = wb.create_sheet("Otra")
    ws2.append(["nota", "valor"])
    ws2.append(["CONSULTA_EXTERNA", "bachillerato"])
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    crear_excel_prueba("datos_originales.xlsx")
    print("Creado datos_originales.xlsx")
