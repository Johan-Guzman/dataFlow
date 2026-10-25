import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import automatizacion_excel as ae
from crear_excel_prueba import ENCABEZADOS, FILAS, crear_excel_prueba

CONFIG = Path(__file__).parent / "config_prueba.json"


def sha(ruta):
    return hashlib.sha256(Path(ruta).read_bytes()).hexdigest()


@pytest.fixture
def ejecucion(tmp_path):
    entrada = crear_excel_prueba(tmp_path / "datos_originales.xlsx")
    hash_antes = sha(entrada)
    salida = tmp_path / "datos_corregidos.xlsx"
    cfg = ae.cargar_configuracion_json(CONFIG)
    resultado = ae.procesar_excel(entrada, salida, cfg)
    return {
        "entrada": entrada, "salida": salida, "hash_antes": hash_antes,
        "resultado": resultado,
        "orig": load_workbook(entrada)["Datos"],
        "out": load_workbook(salida)["Datos"],
        "wb": load_workbook(salida),
    }


def col(ws, nombre):
    for c in ws[1]:
        if c.value == nombre:
            return c.column
    raise KeyError(nombre)


def val(ws, nombre, fila):
    return ws.cell(row=fila, column=col(ws, nombre)).value


# filas: id n -> fila n+1
def test_01_bachillerato_oficial(ejecucion):
    assert val(ejecucion["out"], "nivel educativo", 2) == "Bachillerato"
    assert val(ejecucion["out"], "nivel educativo", 7) == "Bachillerato"  # con espacio final


def test_02_tecnico_oficial(ejecucion):
    assert val(ejecucion["out"], "nivel educativo", 3) == "Técnico"


def test_03_cc_a_denominacion_oficial(ejecucion):
    assert val(ejecucion["out"], "tipo de documento", 2) == "Cédula de ciudadanía"
    # abreviatura no configurada: se conserva y se reporta
    assert val(ejecucion["out"], "tipo de documento", 7) == "RC"
    assert any(r["Valor original"] == "RC" and r["Resultado"] == "INVALIDO"
               for r in ejecucion["resultado"]["reporte"])


def test_04_ingenieria_sin_tilde(ejecucion):
    assert val(ejecucion["out"], "programa de interés", 2) == "Ingeniería de Sistemas"


def test_05_consulta_externa_title_case(ejecucion):
    out = ejecucion["out"]
    assert val(out, "categoria_servicio", 2) == "Consulta Externa"
    assert val(out, "categoria_servicio", 3) == "Medicina General"
    assert val(out, "categoria_servicio", 4) == "Atencion Integral En Salud"
    assert val(out, "categoria_servicio", 5) == "Servicios De Urgencias"
    assert val(out, "categoria_servicio", 6) == "Odontologia"
    assert val(out, "categoria_servicio", 9) == "Apoyo Psicosocial"
    assert val(out, "categoria_servicio", 10) == "Apoyo Psicosocial"


def test_06_telefono_se_separa(ejecucion):
    out = ejecucion["out"]
    assert val(out, "telefono", 2) == "3001234567"
    assert val(out, "Codigo_Pais", 2) == "+57"
    assert val(out, "telefono", 3) == "312 555 0182"  # conserva espacios internos
    assert val(out, "telefono", 5) == "3155551234"
    assert isinstance(val(out, "telefono", 2), str)


def test_06b_casos_especiales_telefono(ejecucion):
    out = ejecucion["out"]
    # solo "+57": se conserva y se reporta
    assert val(out, "telefono", 6) == "+57"
    assert val(out, "Codigo_Pais", 6) is None
    # Codigo_Pais distinto: no se sobrescribe, teléfono intacto
    assert val(out, "telefono", 7) == "+573001234567"
    assert val(out, "Codigo_Pais", 7) == "+1"
    # cero inicial conservado
    assert val(out, "telefono", 10) == "0123456789"
    # número entero: no se toca ni se le agrega indicativo
    assert val(out, "telefono", 9) == 3105551234
    assert val(out, "Codigo_Pais", 9) is None
    motivos = [r for r in ejecucion["resultado"]["reporte"] if r["Resultado"] == "REQUIERE_REVISION"]
    assert any(r["Valor original"] == "+57" for r in motivos)
    assert any(r["Valor original"] == "+1" for r in motivos)


def test_07_telefono_sin_indicativo(ejecucion):
    out = ejecucion["out"]
    assert val(out, "telefono", 4) == "3009876543"
    assert val(out, "Codigo_Pais", 4) is None


def test_08_09_periodos_reconocidos(ejecucion):
    out = ejecucion["out"]
    assert val(out, "periodo", 2) == "202710 - Primer Semestre 2027"
    assert val(out, "periodo", 3) == "202720 - Segundo Semestre 2027"
    assert val(out, "periodo", 5) == "202810 - Primer Semestre 2028"
    assert val(out, "periodo", 7) == "202710 - Primer Semestre 2027"  # ya formateado
    assert val(out, "periodo", 9) == "202720 - Segundo Semestre 2027"  # numérico


def test_10_periodo_no_configurado_o_ambiguo(ejecucion):
    out = ejecucion["out"]
    assert val(out, "periodo", 4) == "202830"
    assert val(out, "periodo", 6) == "periodo pendiente"
    # descripción que contradice el código: se conserva
    assert val(out, "periodo", 10) == "202710 - Segundo Semestre 2027"
    rep = ejecucion["resultado"]["reporte"]
    for original in ("202830", "periodo pendiente", "202710 - Segundo Semestre 2027"):
        assert any(r["Valor original"] == original and r["Resultado"] == "REQUIERE_REVISION"
                   for r in rep)


def test_11_programa_no_catalogado(ejecucion):
    assert val(ejecucion["out"], "programa de interés", 4) == "Programa no catalogado"
    assert any(r["Valor original"] == "Programa no catalogado" and r["Resultado"] == "INVALIDO"
               for r in ejecucion["resultado"]["reporte"])


def test_12_vacios_no_se_completan(ejecucion):
    out = ejecucion["out"]
    for nombre in ENCABEZADOS[2:9]:
        assert val(out, nombre, 8) is None, nombre
    assert not any(r["Fila"] == 8 for r in ejecucion["resultado"]["reporte"])


def test_13_columnas_ajenas_intactas(ejecucion):
    o, n = ejecucion["orig"], ejecucion["out"]
    for nombre in ("id", "nombre", "observaciones", "cantidad", "total"):
        for fila in range(1, o.max_row + 1):
            assert val(n, nombre, fila) == val(o, nombre, fila), (nombre, fila)
    assert n["K2"].number_format == "0.00"
    # formulas conservadas
    assert val(n, "total", 2) == "=K2*2"
    # otra hoja sin columnas configuradas: igual
    otra = ejecucion["wb"]["Otra"]
    assert [c.value for c in otra[2]] == ["CONSULTA_EXTERNA", "bachillerato"]


def test_14_no_elimina_filas_ni_hojas(ejecucion):
    o, n = ejecucion["orig"], ejecucion["out"]
    assert n.max_row == o.max_row == len(FILAS) + 1
    assert [c.value for c in n[1]][: len(ENCABEZADOS)] == ENCABEZADOS
    assert [val(n, "id", f) for f in range(2, n.max_row + 1)] == [f[0] for f in FILAS]
    assert {"Datos", "Otra", "Reporte_validacion", "Resumen_validacion"} <= set(ejecucion["wb"].sheetnames)


def test_original_no_se_modifica(ejecucion):
    assert sha(ejecucion["entrada"]) == ejecucion["hash_antes"]


def test_reporte_estructura(ejecucion):
    hoja = ejecucion["wb"]["Reporte_validacion"]
    assert [c.value for c in hoja[1]] == [
        "Hoja", "Fila", "Columna", "Valor original", "Valor final", "Resultado", "Motivo"]
    estados = {r["Resultado"] for r in ejecucion["resultado"]["reporte"]}
    assert {"CORREGIDO", "INVALIDO", "REQUIERE_REVISION"} <= estados
    assert ejecucion["resultado"]["resumen"]["VALIDO"] > 0  # contados en el resumen


def test_sin_catalogo_resume_en_una_fila(tmp_path):
    entrada = crear_excel_prueba(tmp_path / "in.xlsx")
    r = ae.procesar_excel(entrada, tmp_path / "out.xlsx", ae.Configuracion())
    sc = [x for x in r["reporte"] if x["Resultado"] == "SIN_CATALOGO"]
    assert len(sc) == 3  # una fila por columna, no una por celda
    out = load_workbook(tmp_path / "out.xlsx")["Datos"]
    assert val(out, "nivel educativo", 2) == "bachillerato"  # sin catálogo: no se corrige
    assert val(out, "tipo de documento", 2) == "CC"          # sin equivalencia: no se asume


def test_no_pisa_original_con_misma_ruta(tmp_path):
    entrada = crear_excel_prueba(tmp_path / "in.xlsx")
    with pytest.raises(ValueError):
        ae.procesar_excel(entrada, entrada, ae.Configuracion())


def test_hoja_reglas_no_se_modifica_y_catalogos_opcionales(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["nivel educativo"])
    ws.append(["bachillerato"])
    reglas = wb.create_sheet("Reglas de validación")
    reglas.append(["nivel educativo"])
    reglas.append(["Bachillerato"])
    reglas.append(["tecnico"])  # la hoja de reglas no debe normalizarse
    wb.save(tmp_path / "in.xlsx")
    cfg = ae.Configuracion(leer_catalogos_desde_hoja=True)
    ae.procesar_excel(tmp_path / "in.xlsx", tmp_path / "out.xlsx", cfg)
    out = load_workbook(tmp_path / "out.xlsx")
    assert out["Datos"]["A2"].value == "Bachillerato"
    assert out["Reglas de validación"]["A3"].value == "tecnico"


def test_funciones_unitarias():
    assert ae.convertir_title_case("CONSULTA__EXTERNA") == "Consulta Externa"
    assert ae.convertir_title_case(" a_b ") == "A B"
    assert ae.normalizar_telefono("+57")[2] is False
    assert ae.normalizar_periodo("202830", {"10": "P"})[1] == "REQUIERE_REVISION"
    assert ae.normalizar_periodo("199910", {"10": "P"})[1] == "REQUIERE_REVISION"  # año fuera de rango
    # la ñ no se confunde con n
    assert ae.validar_valor("Diseno", ["Diseño"])[1] == "INVALIDO"


def test_catalogo_ambiguo_no_se_corrige():
    valor, estado, _ = ae.validar_valor("bachillerato", ["Bachillerato", "BACHILLERATO"])
    assert (valor, estado) == ("bachillerato", "REQUIERE_REVISION")
